import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from cb_policy_tone import gold_sample as gs


def make_pool(per_era: int = 1500, seed: int = 0, hawkish_share: float = 0.2) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    rows = []
    for era in gs.QUOTAS:
        for i in range(per_era):
            tone = rng.choice(gs.TONES, p=[hawkish_share, 1 - hawkish_share - 0.2, 0.2])
            topic = rng.choice([*gs.TOPICS, "(none)"])
            doc_id = f"{era}-{int(rng.integers(0, 60))}"
            text = f"Sentence {era} number {i} about {topic}."
            rows.append(
                {
                    "sentence_id": f"{doc_id}-{i:04d}", "doc_id": doc_id, "meeting_no": 5,
                    "text": text, "norm_text": gs.normalise(text), "section_heading": "",
                    "tone_weak": tone, "topic_weak": topic, "era": era,
                }
            )  # fmt: skip
    return pd.DataFrame(rows)


@pytest.fixture(scope="module")
def sample():
    return gs.draw_sample(make_pool())


# ---- the draw ----


def test_each_era_gets_exactly_its_quota(sample):
    assert sample["era"].value_counts().to_dict() == gs.QUOTAS
    assert len(sample) == 400


def test_forty_percent_are_a_simple_random_draw(sample):
    assert (sample["draw"] == "random").sum() == 36 + 36 + 32 + 56


def test_the_stratified_part_balances_the_tone_classes(sample):
    stratified = sample[sample["draw"] == "stratified"]
    for _, era in stratified.groupby("era"):
        counts = era["tone_weak"].value_counts()
        assert counts.max() - counts.min() <= 1


def test_no_repeated_text_and_a_cap_per_statement(sample):
    assert not sample["norm_text"].duplicated().any()
    assert sample["doc_id"].value_counts().max() <= gs.MAX_PER_DOC


def test_a_quarter_is_dev_and_the_out_of_regime_test_set_is_big_enough(sample):
    assert abs((sample["split"] == "dev").sum() - 100) <= 5
    assert ((sample["era"] == "2022-2026") & (sample["split"] == "test")).sum() >= 100


def test_gold_ids_are_unique_and_sequential(sample):
    assert sample["gold_id"].tolist() == [f"G{i:03d}" for i in range(1, 401)]


def test_the_same_seed_gives_the_same_sample_and_another_seed_does_not(sample):
    pool = make_pool()
    again = gs.draw_sample(pool)
    other = gs.draw_sample(pool, seed=7)
    assert again["sentence_id"].tolist() == sample["sentence_id"].tolist()
    assert other["sentence_id"].tolist() != sample["sentence_id"].tolist()


def test_a_scarce_tone_class_is_topped_up_from_the_rest():
    pool = make_pool(hawkish_share=0.0)  # no hawkish sentences at all
    result = gs.draw_sample(pool)
    assert result["era"].value_counts().to_dict() == gs.QUOTAS
    assert (result["tone_weak"] == "hawkish").sum() == 0


def test_the_key_has_no_text_and_the_sheet_no_lexicon_labels(sample, tmp_path):
    key = gs.make_key(sample)
    assert "text" not in key.columns
    assert {"weak_tone", "weak_topic", "split"} <= set(key.columns)

    path = tmp_path / "sheet.xlsx"
    gs.write_annotation(sample, path)
    sheet = load_workbook(path)["labels"]
    assert [c.value for c in sheet[1]] == gs.HEADERS
    assert all(
        sheet.cell(row=r, column=c).value is None for r in range(2, 402) for c in (4, 5, 6, 7)
    )
    assert not any("weak" in str(c.value) for c in sheet[1])


# ---- the sheet ----


def test_the_sheet_has_dropdowns_for_tone_topic_and_unsure(sample, tmp_path):
    path = tmp_path / "sheet.xlsx"
    gs.write_annotation(sample, path)
    workbook = load_workbook(path)
    lists = {str(v.sqref): v.formula1 for v in workbook["labels"].data_validations.dataValidation}
    assert lists["D2:D401"] == '"hawkish,neutral,dovish"'
    assert lists["E2:E401"].startswith('"inflation,growth')
    assert lists["F2:F401"] == '"yes"'
    assert "rules" in workbook.sheetnames


def test_check_labels_reports_progress_and_invalid_entries(sample, tmp_path):
    path = tmp_path / "sheet.xlsx"
    gs.write_annotation(sample, path)
    workbook = load_workbook(path)
    sheet = workbook["labels"]
    sheet["D2"], sheet["E2"] = "hawkish", "inflation"
    sheet["D3"], sheet["E3"], sheet["F3"] = "dovish", "growth", "yes"
    sheet["D4"], sheet["E4"] = "hawk", "inflation"  # a typo that bypassed the dropdown
    workbook.save(path)

    result = gs.check_labels(path)
    assert result["rows"] == 400
    assert result["tone_done"] == 3
    assert result["unsure"] == 1
    assert result["invalid_tone"] == [sample["gold_id"].iloc[2]]


# ---- the pool ----


def test_build_pool_uses_the_letterhead_date_when_the_meeting_has_no_date():
    sentences = pd.DataFrame(
        {
            "sentence_id": ["a-1", "b-1", "c-1"],
            "doc_id": [1, 2, 3],
            "meeting_end": ["2010-05-01", None, "1999-01-01"],
            "text": ["First sentence here.", "Second sentence here.", "Third sentence here."],
        }
    )
    weak = pd.DataFrame(
        {
            "sentence_id": ["a-1", "b-1", "c-1"],
            "labellable": [True, True, True],
            "tone_weak": ["neutral"] * 3,
            "topic_weak": ["inflation", None, "growth"],
        }
    )
    statements = pd.DataFrame({"doc_id": [1, 2, 3], "letterhead_date": [None, "2025-07-18", None]})
    pool = gs.build_pool(sentences, weak, statements).set_index("sentence_id")
    assert pool.loc["a-1", "era"] == "2010-2016"
    assert pool.loc["b-1", "era"] == "2022-2026"
    assert "c-1" not in pool.index  # 1999 is outside every era
    assert pool.loc["b-1", "topic_weak"] == "(none)"


def test_drawing_again_refuses_to_overwrite_existing_labels(tmp_path, monkeypatch, capsys):
    sheet = tmp_path / "gold_annotation.xlsx"
    sheet.write_text("hand-labelled work")
    monkeypatch.setattr(gs, "ANNOTATION", sheet)
    monkeypatch.setattr(gs, "KEY", tmp_path / "key.csv")
    gs.draw()
    assert sheet.read_text() == "hand-labelled work"
    assert "Refusing to overwrite" in capsys.readouterr().out
