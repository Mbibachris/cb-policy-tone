import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from cb_policy_tone import evaluate_gold as eg
from cb_policy_tone import gold_sample as gs


def make_sample(n: int = 60) -> pd.DataFrame:
    rows = []
    for i in range(n):
        rows.append(
            {
                "gold_id": f"G{i + 1:03d}", "sentence_id": f"s-{i}", "doc_id": i, "meeting_no": 5,
                "era": "2022-2026" if i % 2 else "2010-2016", "draw": "random" if i % 3 == 0 else "stratified",
                "split": "dev" if i < 15 else "test", "text": f"Sentence {i}.", "section_heading": "",
                "tone_weak": "neutral", "topic_weak": "growth", "norm_text": f"sentence {i}",
            }
        )  # fmt: skip
    return pd.DataFrame(rows)


@pytest.fixture
def labelled_sheet(tmp_path):
    sample = make_sample()
    path = tmp_path / "gold.xlsx"
    gs.write_annotation(sample, path)
    wb = load_workbook(path)
    ws = wb["labels"]
    for i in range(len(sample)):
        ws.cell(row=i + 2, column=4, value=["hawkish", "neutral", "dovish"][i % 3])
        ws.cell(row=i + 2, column=5, value="growth" if i % 2 else "inflation")
    ws["D5"] = "hawk"  # a typo: this row must be ignored
    ws["F6"] = "yes"
    wb.save(path)
    return path, sample


def test_only_rows_with_valid_tone_and_topic_are_loaded(labelled_sheet):
    path, sample = labelled_sheet
    gold = eg.load_gold(path)
    assert len(gold) == len(sample) - 1
    assert "G004" not in set(gold["gold_id"])
    assert gold["unsure"].sum() == 1


def test_the_model_label_is_the_most_probable_class(labelled_sheet):
    path, sample = labelled_sheet
    gold = eg.load_gold(path)
    key = gs.make_key(sample.assign(tone_weak="neutral"))
    weak = pd.DataFrame({"sentence_id": sample["sentence_id"], "hawk_rules": "", "dove_rules": ""})
    probs = pd.DataFrame(
        {"sentence_id": sample["sentence_id"], "p_dovish": 0.1, "p_neutral": 0.2, "p_hawkish": 0.7}
    )
    df = eg.assemble(gold, key, weak, probs)
    assert set(df["model_tone"]) == {"hawkish"}
    assert eg.assemble(gold, key, weak, None)["model_tone"].isna().all()


def test_score_matches_a_hand_calculation():
    truth = pd.Series(["hawkish", "hawkish", "neutral", "dovish"])
    pred = pd.Series(["hawkish", "neutral", "neutral", "dovish"])
    result = eg.score(truth, pred)
    assert result["accuracy"] == 0.75
    assert result["f1"] == {
        "hawkish": pytest.approx(2 / 3),
        "neutral": pytest.approx(2 / 3),
        "dovish": 1.0,
    }
    assert result["macro_f1"] == pytest.approx((2 / 3 + 2 / 3 + 1) / 3)


def test_bootstrap_interval_contains_the_point_estimate_and_difference_is_reported():
    rng = np.random.default_rng(1)
    truth = pd.Series(rng.choice(eg.ORDER, 300))
    good = truth.where(rng.random(300) < 0.8, pd.Series(rng.choice(eg.ORDER, 300)))
    poor = pd.Series(rng.choice(eg.ORDER, 300))
    out = eg.bootstrap(truth, {"lexicon": poor, "model": good}, n=200)
    point = eg.score(truth, good)["macro_f1"]
    assert out["model"][0] <= point <= out["model"][1]
    assert out["model minus lexicon"][0] > 0  # the clearly better system wins


def test_dev_errors_export_never_contains_a_test_sentence(labelled_sheet):
    path, sample = labelled_sheet
    gold = eg.load_gold(path)
    key = gs.make_key(sample.assign(tone_weak="neutral"))
    weak = pd.DataFrame(
        {"sentence_id": sample["sentence_id"], "hawk_rules": "T01", "dove_rules": ""}
    )
    df = eg.assemble(gold, key, weak, None)
    errors = eg.dev_errors(df)
    assert len(errors) > 0
    assert set(errors["gold_id"]) <= set(sample.loc[sample["split"] == "dev", "gold_id"])


def test_subsets_split_dev_test_and_the_random_draw(labelled_sheet):
    path, sample = labelled_sheet
    gold = eg.load_gold(path)
    df = eg.assemble(
        gold,
        gs.make_key(sample),
        pd.DataFrame({"sentence_id": sample["sentence_id"], "hawk_rules": "", "dove_rules": ""}),
        None,
    )
    parts = eg.subsets(df)
    assert set(parts["gold-dev"]["split"]) == {"dev"}
    assert set(parts["gold-test"]["split"]) == {"test"}
    assert set(parts["gold-test 2022-2026"]["era"]) == {"2022-2026"}
    assert set(parts["random draw only (dev+test)"]["draw"]) == {"random"}


def test_the_report_runs_with_and_without_a_model(labelled_sheet, capsys):
    path, sample = labelled_sheet
    gold = eg.load_gold(path)
    key = gs.make_key(sample)
    weak = pd.DataFrame({"sentence_id": sample["sentence_id"], "hawk_rules": "", "dove_rules": ""})
    probs = pd.DataFrame(
        {"sentence_id": sample["sentence_id"], "p_dovish": 0.2, "p_neutral": 0.5, "p_hawkish": 0.3}
    )
    eg.report(eg.assemble(gold, key, weak, None))
    assert "lexicon" in capsys.readouterr().out
    eg.report(eg.assemble(gold, key, weak, probs))
    assert "model minus lexicon" in capsys.readouterr().out
