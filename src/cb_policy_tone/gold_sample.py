"""Draw the blind gold-standard sample and write the sheet for hand labelling.

Reads : data/processed/sentences.csv, weak_labels.csv and statements.csv
Writes: data/labels/gold_annotation.xlsx      the sheet to label (blind: no lexicon labels)
        data/labels/gold_key_DO_NOT_OPEN.csv  ids, era, draw type, dev/test split, lexicon labels

Commands:
    python -m cb_policy_tone.gold_sample           draw the sample (refuses to overwrite)
    python -m cb_policy_tone.gold_sample check     report progress and invalid labels

Design (see the codebook, section 6):
- 400 sentences from the labellable pool, with fixed quotas per era so that 2022-2026 is well
  represented for the out-of-regime test.
- 40 percent are a simple random draw (unbiased, for honest baseline numbers). 60 percent are
  stratified by the lexicon's tone class, so that hawkish and dovish sentences are not drowned
  out by neutral ones. The key records which is which.
- No repeated sentence text, and at most 8 sentences per statement.
- About 25 percent form gold-dev (for tuning); about 75 percent form gold-test (never tuned on).
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SENTENCES = Path("data/processed/sentences.csv")
WEAK = Path("data/processed/weak_labels.csv")
STATEMENTS = Path("data/processed/statements.csv")
ANNOTATION = Path("data/labels/gold_annotation.xlsx")
KEY = Path("data/labels/gold_key_DO_NOT_OPEN.csv")

QUOTAS = {"2002-2009": 90, "2010-2016": 90, "2017-2021": 80, "2022-2026": 140}
ERA_YEARS = {
    "2002-2009": (2002, 2009),
    "2010-2016": (2010, 2016),
    "2017-2021": (2017, 2021),
    "2022-2026": (2022, 2026),
}
RANDOM_SHARE = 0.4
DEV_SHARE = 0.25
MAX_PER_DOC = 8
SEED = 42

TONES = ["hawkish", "neutral", "dovish"]
TOPICS = ["inflation", "growth", "external", "fiscal", "financial", "global", "policy", "other"]
HEADERS = ["gold_id", "section_heading", "sentence", "tone", "topic", "unsure", "notes"]

RULES_SHEET = [
    ("TONE", "Judge the policy implication, not whether the news is good or bad."),
    ("", "If this were the only sentence a reader saw, which way would they expect the next rate move?"),
    ("hawkish", "points to tighter policy"),
    ("dovish", "points to looser policy"),
    ("neutral", "points neither way, or two signals of similar weight offset each other"),
    ("", ""),
    ("DIRECTION", "hawkish | dovish"),
    ("Inflation", "rising, above the target band, upside risks | falling, below or back in the band, downside risks"),
    ("Growth (call 2)", "strong growth is NEUTRAL unless demand is above potential | weak growth, downside risks to growth"),
    ("Cedi (call 1)", "depreciation, pressure on the cedi | appreciation, stability"),
    ("Reserves, exports (calls 3, 6)", "(no hawkish rule) | stronger reserves, higher export receipts, higher gold or cocoa prices"),
    ("Fiscal (call 4)", "deficit overshoot, revenue shortfall, more government or central-bank borrowing | consolidation, deficit on target"),
    ("Credit (call 5)", "fast money or credit growth ONLY if framed as liquidity or inflation pressure | weak or contracting credit"),
    ("Oil (call 6)", "higher oil prices | lower oil prices"),
    ("Yields (call 7)", "NEUTRAL: market outcomes, not signals"),
    ("Bank soundness (call 8)", "NEUTRAL unless the sentence calls for caution or support"),
    ("Global", "rising global inflation, tightening by other central banks, strong dollar | global slowdown, easing abroad"),
    ("Guidance", "'remain vigilant', 'further tightening' | 'room to ease', 'accommodative', 'support growth'"),
    ("", ""),
    ("TOPIC", "One primary topic. If the sentence describes a Bank action or guidance: policy."),
    ("", "If its subject is another country or a world aggregate: global. Otherwise the topic of the grammatical subject."),
    ("inflation", "prices, expectations, the target band"),
    ("growth", "GDP, activity, PMI, confidence, sector output, employment"),
    ("external", "the cedi, balance of payments, trade, reserves, remittances"),
    ("fiscal", "budget, revenue, expenditure, deficit, public debt, government borrowing"),
    ("financial", "money and credit, interest rates and yields, bank soundness, markets"),
    ("global", "the world economy, other countries, other central banks, global commodity prices"),
    ("policy", "the Bank's stance, guidance, measures such as reserve requirements"),
    ("other", "greetings, procedure, scheduling, nothing economic"),
    ("", ""),
    ("DOUBT", "Between hawkish and neutral: choose neutral and set 'unsure' to yes. Do not guess."),
    ("RULE", "Do not open the key file. Do not look at the lexicon labels until every sentence is done."),
]  # fmt: skip


def era_of(year: float) -> str:
    if pd.isna(year):
        return ""
    for era, (low, high) in ERA_YEARS.items():
        if low <= year <= high:
            return era
    return ""


def normalise(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def build_pool(
    sentences: pd.DataFrame, weak: pd.DataFrame, statements: pd.DataFrame
) -> pd.DataFrame:
    """Labellable sentences with an era. Sentences with no meeting date use the letterhead date."""
    pool = sentences.merge(weak[weak["labellable"]], on="sentence_id")
    letter = statements[["doc_id", "letterhead_date"]].drop_duplicates("doc_id")
    pool = pool.merge(letter, on="doc_id", how="left")
    date = pd.to_datetime(pool["meeting_end"]).fillna(pd.to_datetime(pool["letterhead_date"]))
    pool["year"] = date.dt.year
    pool["era"] = pool["year"].map(era_of)
    pool["norm_text"] = pool["text"].map(normalise)
    pool["topic_weak"] = pool["topic_weak"].fillna("").replace("", "(none)")
    return pool[pool["era"] != ""].reset_index(drop=True)


def split_evenly(total: int, parts: int) -> list[int]:
    base, extra = divmod(total, parts)
    return [base + (1 if i < extra else 0) for i in range(parts)]


def round_robin(candidates: pd.DataFrame, rng: np.random.Generator) -> list[pd.Series]:
    """Rows interleaved across weak topics, each topic shuffled, so topics are spread out."""
    groups = [
        group.sample(frac=1, random_state=int(rng.integers(1_000_000_000)))
        for _, group in candidates.groupby("topic_weak")
    ]
    order = []
    for i in range(max((len(g) for g in groups), default=0)):
        for group in groups:
            if i < len(group):
                order.append(group.iloc[i])
    return order


class Chooser:
    """Adds rows while keeping texts unique and at most MAX_PER_DOC sentences per statement."""

    def __init__(self) -> None:
        self.texts: set[str] = set()
        self.per_doc: Counter = Counter()
        self.rows: list[dict] = []

    def add(self, row: pd.Series, draw: str) -> bool:
        if row["norm_text"] in self.texts or self.per_doc[row["doc_id"]] >= MAX_PER_DOC:
            return False
        self.texts.add(row["norm_text"])
        self.per_doc[row["doc_id"]] += 1
        self.rows.append({**row.to_dict(), "draw": draw})
        return True


def draw_sample(
    pool: pd.DataFrame, seed: int = SEED, quotas: dict[str, int] = QUOTAS
) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    chooser = Chooser()
    for era, quota in quotas.items():
        era_pool = pool[pool["era"] == era]
        n_random = round(RANDOM_SHARE * quota)
        added = 0
        for _, row in era_pool.sample(
            frac=1, random_state=int(rng.integers(1_000_000_000))
        ).iterrows():
            if added >= n_random:
                break
            added += chooser.add(row, "random")

        remaining = quota - added
        taken = {r["sentence_id"] for r in chooser.rows}
        shortfall = 0
        for tone, share in zip(TONES, split_evenly(remaining, len(TONES)), strict=True):
            candidates = era_pool[
                (era_pool["tone_weak"] == tone) & ~era_pool["sentence_id"].isin(taken)
            ]
            got = 0
            for row in round_robin(candidates, rng):
                if got >= share:
                    break
                got += chooser.add(row, "stratified")
            shortfall += share - got
        if shortfall:  # a class ran short: top up from whatever is left in the era
            taken = {r["sentence_id"] for r in chooser.rows}
            rest = era_pool[~era_pool["sentence_id"].isin(taken)]
            for row in round_robin(rest, rng):
                if shortfall <= 0:
                    break
                shortfall -= chooser.add(row, "stratified")

    sample = pd.DataFrame(chooser.rows)
    sample["split"] = "test"
    for _, cell in sample.groupby(["era", "draw"]):
        shuffled = cell.sample(frac=1, random_state=int(rng.integers(1_000_000_000)))
        sample.loc[shuffled.index[: round(DEV_SHARE * len(cell))], "split"] = "dev"
    sample = sample.sample(frac=1, random_state=int(rng.integers(1_000_000_000))).reset_index(
        drop=True
    )
    sample.insert(0, "gold_id", [f"G{i:03d}" for i in range(1, len(sample) + 1)])
    return sample


def make_key(sample: pd.DataFrame) -> pd.DataFrame:
    key = sample[["gold_id", "sentence_id", "doc_id", "meeting_no", "era", "draw", "split"]].copy()
    key["weak_tone"] = sample["tone_weak"]
    key["weak_topic"] = sample["topic_weak"]
    return key


def write_annotation(sample: pd.DataFrame, path: Path = ANNOTATION) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "labels"
    ws.append(HEADERS)
    for r in sample.itertuples():
        ws.append(
            [
                r.gold_id,
                r.section_heading if isinstance(r.section_heading, str) else "",
                r.text,
                None,
                None,
                None,
                None,
            ]
        )
    last = len(sample) + 1

    for col, width in zip("ABCDEFG", (8, 28, 110, 12, 12, 8, 30), strict=True):
        ws.column_dimensions[col].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in ws.iter_rows(min_row=2, max_row=last, min_col=2, max_col=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"

    for column, options in (("D", TONES), ("E", TOPICS), ("F", ["yes"])):
        validation = DataValidation(
            type="list", formula1='"' + ",".join(options) + '"', allow_blank=True
        )
        validation.showErrorMessage = True
        validation.errorTitle = "Not allowed"
        validation.error = "Pick a value from the list."
        ws.add_data_validation(validation)
        validation.add(f"{column}2:{column}{last}")

    rules = wb.create_sheet("rules")
    for left, right in RULES_SHEET:
        rules.append([left, right])
    rules.column_dimensions["A"].width = 30
    rules.column_dimensions["B"].width = 140
    for row in rules.iter_rows():
        row[0].font = Font(bold=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def check_labels(path: Path = ANNOTATION) -> dict:
    """Progress and problems in a (partly) labelled sheet."""
    wb = load_workbook(path, read_only=True)
    rows = list(wb["labels"].iter_rows(min_row=2, values_only=True))
    frame = pd.DataFrame(rows, columns=HEADERS)
    tone = frame["tone"].fillna("")
    topic = frame["topic"].fillna("")
    return {
        "rows": len(frame),
        "tone_done": int((tone != "").sum()),
        "topic_done": int((topic != "").sum()),
        "invalid_tone": frame.loc[(tone != "") & ~tone.isin(TONES), "gold_id"].tolist(),
        "invalid_topic": frame.loc[(topic != "") & ~topic.isin(TOPICS), "gold_id"].tolist(),
        "unsure": int((frame["unsure"].fillna("") == "yes").sum()),
        "tone_counts": tone[tone != ""].value_counts().to_dict(),
    }


def draw(force: bool = False) -> None:
    if (ANNOTATION.exists() or KEY.exists()) and not force:
        print(f"{ANNOTATION} or {KEY} already exists. Refusing to overwrite: labels could be lost.")
        print("Use --force only if you are sure you want to start again.")
        return
    sentences = pd.read_csv(SENTENCES)
    sentences["section_heading"] = sentences["section_heading"].fillna("")
    pool = build_pool(sentences, pd.read_csv(WEAK), pd.read_csv(STATEMENTS))
    sample = draw_sample(pool)
    write_annotation(sample)
    key = make_key(sample)
    key.to_csv(KEY, index=False)

    print(f"Pool: {len(pool)} labellable sentences. Sample: {len(sample)}.")
    print(sample.groupby(["era", "draw"]).size().unstack(fill_value=0).to_string())
    print("\nSplit:", sample["split"].value_counts().to_dict())
    print(
        "2022-2026 in gold-test:",
        int(((sample["era"] == "2022-2026") & (sample["split"] == "test")).sum()),
    )
    print(f"Statements represented: {sample['doc_id'].nunique()}")
    print(f"\nSheet to label: {ANNOTATION}\nKey (do not open): {KEY}")


def report_progress() -> None:
    sys.stdout.reconfigure(errors="replace")
    result = check_labels()
    print(
        f"Rows: {result['rows']} | tone labelled: {result['tone_done']} | topic labelled: {result['topic_done']}"
    )
    print(f"Marked unsure: {result['unsure']} | tone counts so far: {result['tone_counts']}")
    for name in ("invalid_tone", "invalid_topic"):
        if result[name]:
            print(f"PROBLEM, {name}: {result[name][:20]}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawTextHelpFormatter
    )
    parser.add_argument("command", nargs="?", default="draw", choices=["draw", "check"])
    parser.add_argument("--force", action="store_true", help="overwrite an existing sheet and key")
    args = parser.parse_args()
    if args.command == "check":
        report_progress()
    else:
        draw(force=args.force)


if __name__ == "__main__":
    main()
