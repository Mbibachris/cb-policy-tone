"""Score the lexicon and the fine-tuned model against the hand-labelled gold set.

Reads : data/labels/gold_annotation.xlsx        your labels (partial labelling is fine)
        data/labels/gold_key_DO_NOT_OPEN.csv    split, draw type and lexicon labels (ids only)
        data/processed/weak_labels.csv          lexicon labels and the rules that fired
        data/processed/predictions_tone.csv     model probabilities (optional)
        data/processed/predictions_fed.csv      Fed-trained model, zero-shot, by gold_id (optional)
Writes: data/labels/gold_dev_errors.csv         DEV sentences where the lexicon disagrees with you

Discipline:
- The lexicon is tuned on gold-DEV only. This script exports errors for the dev split only and
  never prints or exports a gold-test sentence.
- gold-TEST is scored once the rules are frozen, and the numbers go in the README as they are.
- 'Random draw only' scores use the unbiased part of the sample, so they are the honest ones
  for prevalence-dependent numbers such as accuracy.
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.metrics import accuracy_score, confusion_matrix, f1_score

from cb_policy_tone.gold_sample import ANNOTATION, HEADERS, KEY, TONES, TOPICS

WEAK = Path("data/processed/weak_labels.csv")
PREDICTIONS = Path("data/processed/predictions_tone.csv")
PREDICTIONS_FED = Path("data/processed/predictions_fed.csv")
DEV_ERRORS = Path("data/labels/gold_dev_errors.csv")
MODEL_LABELS = ["dovish", "neutral", "hawkish"]  # the order of the model's probability columns
ORDER = ["hawkish", "neutral", "dovish"]
N_BOOT = 1000
SEED = 0


def load_gold(path: Path = ANNOTATION) -> pd.DataFrame:
    """Rows with a valid tone AND topic. Partly labelled sheets are fine."""
    wb = load_workbook(path, read_only=True)
    frame = pd.DataFrame(list(wb["labels"].iter_rows(min_row=2, values_only=True)), columns=HEADERS)
    ok = frame["tone"].isin(TONES) & frame["topic"].isin(TOPICS)
    frame = frame[ok].rename(columns={"tone": "gold_tone", "topic": "gold_topic"})
    frame["unsure"] = frame["unsure"].fillna("") == "yes"
    return frame.reset_index(drop=True)


def assemble(
    gold: pd.DataFrame,
    key: pd.DataFrame,
    weak: pd.DataFrame,
    predictions: pd.DataFrame | None,
    fed: pd.DataFrame | None = None,
) -> pd.DataFrame:
    df = gold.merge(key, on="gold_id", how="left")
    df = df.merge(weak[["sentence_id", "hawk_rules", "dove_rules"]], on="sentence_id", how="left")
    df["lexicon_tone"] = df["weak_tone"]
    df["lexicon_topic"] = df["weak_topic"]
    df["model_tone"] = np.nan
    df["fed_tone"] = np.nan
    if predictions is not None:
        probs = predictions.set_index("sentence_id")[[f"p_{label}" for label in MODEL_LABELS]]
        best = probs.to_numpy().argmax(axis=1)
        picked = pd.Series([MODEL_LABELS[i] for i in best], index=probs.index)
        df["model_tone"] = df["sentence_id"].map(picked)
    if fed is not None:
        df["fed_tone"] = df["gold_id"].map(fed.set_index("gold_id")["fed_label"])
    return df


def subsets(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    return {
        "gold-dev": df[df["split"] == "dev"],
        "gold-test": df[df["split"] == "test"],
        "gold-test 2022-2026": df[(df["split"] == "test") & (df["era"] == "2022-2026")],
        "random draw only (dev+test)": df[df["draw"] == "random"],
        "everything labelled": df,
    }


def score(truth: pd.Series, pred: pd.Series) -> dict:
    per_class = f1_score(truth, pred, labels=ORDER, average=None, zero_division=0)
    return {
        "n": len(truth),
        "accuracy": accuracy_score(truth, pred) if len(truth) else float("nan"),
        "macro_f1": f1_score(truth, pred, labels=ORDER, average="macro", zero_division=0),
        "f1": dict(zip(ORDER, per_class, strict=True)),
    }


def bootstrap(
    truth: pd.Series, preds: dict[str, pd.Series], n: int = N_BOOT, seed: int = SEED
) -> dict:
    """95% intervals for each system's macro-F1, and for each difference from the first system."""
    rng = np.random.default_rng(seed)
    t = truth.to_numpy()
    p = {name: series.to_numpy() for name, series in preds.items()}
    draws: dict[str, list[float]] = {name: [] for name in p}
    for _ in range(n):
        idx = rng.integers(0, len(t), len(t))
        for name in p:
            draws[name].append(
                f1_score(t[idx], p[name][idx], labels=ORDER, average="macro", zero_division=0)
            )
    out = {name: tuple(np.percentile(v, [2.5, 97.5])) for name, v in draws.items()}
    names = list(p)
    for other in names[1:]:
        diff = np.array(draws[other]) - np.array(draws[names[0]])
        out[f"{other} minus {names[0]}"] = tuple(np.percentile(diff, [2.5, 97.5]))
    return out


def dev_errors(df: pd.DataFrame) -> pd.DataFrame:
    """Dev sentences where the lexicon's tone or topic differs from the hand label."""
    dev = df[df["split"] == "dev"]
    wrong = dev[
        (dev["lexicon_tone"] != dev["gold_tone"]) | (dev["lexicon_topic"] != dev["gold_topic"])
    ]
    cols = [
        "gold_id",
        "sentence",
        "gold_tone",
        "lexicon_tone",
        "model_tone",
        "gold_topic",
        "lexicon_topic",
        "hawk_rules",
        "dove_rules",
        "unsure",
        "notes",
    ]
    return wrong[cols].reset_index(drop=True)


def report(df: pd.DataFrame) -> None:
    sys.stdout.reconfigure(errors="replace")
    has_model = df["model_tone"].notna().any()
    has_fed = df["fed_tone"].notna().any()
    print(
        f"Labelled sentences: {len(df)} (dev {int((df['split'] == 'dev').sum())}, test {int((df['split'] == 'test').sum())})"
    )
    print(
        "Gold tone:",
        df["gold_tone"].value_counts().to_dict(),
        "| marked unsure:",
        int(df["unsure"].sum()),
    )

    for name, part in subsets(df).items():
        if len(part) < 20:
            print(f"\n=== {name}: only {len(part)} labelled; skipped ===")
            continue
        systems = {
            "always neutral": pd.Series("neutral", index=part.index),
            "lexicon": part["lexicon_tone"],
        }
        if has_model:
            systems["model"] = part["model_tone"]
        if has_fed:
            systems["fed zero-shot"] = part["fed_tone"]
        print(f"\n=== {name} (n={len(part)}) ===")
        intervals = bootstrap(
            part["gold_tone"], {k: v for k, v in systems.items() if k != "always neutral"}
        )
        for system, pred in systems.items():
            s = score(part["gold_tone"], pred)
            ci = intervals.get(system)
            ci_text = f" [95% CI {ci[0]:.2f}-{ci[1]:.2f}]" if ci else ""
            f1s = ", ".join(f"{k} {v:.2f}" for k, v in s["f1"].items())
            print(
                f"  {system:15s} accuracy {s['accuracy']:.2f} | macro-F1 {s['macro_f1']:.2f}{ci_text} | F1: {f1s}"
            )
        for other in ("model", "fed zero-shot"):
            if other in systems:
                lo, hi = intervals[f"{other} minus lexicon"]
                print(f"  {other} minus lexicon, macro-F1: 95% CI {lo:+.2f} to {hi:+.2f}")
        topic_ok = (part["lexicon_topic"] == part["gold_topic"]).mean()
        print(f"  lexicon topic accuracy: {topic_ok:.2f}")

    test = df[df["split"] == "test"]
    if len(test) >= 20:
        print("\nConfusion on gold-test (rows = your label, columns = prediction):")
        for system, col, present in (
            ("lexicon", "lexicon_tone", True),
            ("model", "model_tone", has_model),
            ("fed zero-shot", "fed_tone", has_fed),
        ):
            if not present:
                continue
            matrix = confusion_matrix(test["gold_tone"], test[col], labels=ORDER)
            print(f"  {system}:")
            print(
                pd.DataFrame(
                    matrix, index=[f"true {c}" for c in ORDER], columns=[f"pred {c}" for c in ORDER]
                ).to_string()
            )


def main() -> None:
    gold = load_gold()
    key = pd.read_csv(KEY)
    weak = pd.read_csv(WEAK)
    predictions = pd.read_csv(PREDICTIONS) if PREDICTIONS.exists() else None
    fed = pd.read_csv(PREDICTIONS_FED) if PREDICTIONS_FED.exists() else None
    df = assemble(gold, key, weak, predictions, fed)
    report(df)
    errors = dev_errors(df)
    errors.to_csv(DEV_ERRORS, index=False)
    print(f"\nWrote {len(errors)} dev-set disagreements to {DEV_ERRORS}")
    print("(Dev only. Tune the lexicon on these. Never look at gold-test sentences.)")


if __name__ == "__main__":
    main()
